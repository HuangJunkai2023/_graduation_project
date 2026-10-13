import json
from dataclasses import replace
from pathlib import Path
import httpx
import pytest
import pymupdf
from docx import Document
from ocs_qa.config import Settings
from ocs_qa.documents import standards, split_text
from ocs_qa.store import Store
from ocs_qa.api import ModelClient, APIError
from ocs_qa.pipeline import Pipeline, resolve_query
from ocs_qa.evaluation import run_evaluation, retrieval_metrics


TEXT = "接触网零部件表面应光洁，不应有裂纹、疤痕和剥离。检验时应记录表面缺陷的位置。"


@pytest.fixture
def settings(tmp_path):
    return Settings(db_path=tmp_path/"kb.sqlite3",chat_key="test-key",embedding_key="embed-key",
                    embedding_url="https://embedding.example/v1",embedding_model="test-embedding")


def document(tmp_path,name="TB T 2073-2020.docx",text=TEXT):
    path=tmp_path/name
    doc=Document(); doc.add_paragraph(text); doc.save(path)
    return path


@pytest.fixture
def store(settings,tmp_path):
    store=Store(settings)
    assert store.ingest(document(tmp_path),allow_cloud=True)["status"]=="imported"
    return store


class FakeClient:
    def __init__(self, status="answered", quote=TEXT[:20], source=1):
        self.calls=[]; self.status=status; self.quote=quote; self.source=source
    def chat(self,system,prompt):
        self.calls.append(json.loads(prompt))
        return {"status":self.status,"claims":[{"text":"表面应光洁，无裂纹等缺陷。","citations":[{"source":self.source,"quote":self.quote}]}],"missing":""}, {"total_tokens":20}
    def embed(self,texts):
        return [[1.,float("表面" in text)+1.] for text in texts]


@pytest.mark.parametrize("raw,expected",[
    ("TB/T 2073-2020",("TBT2073","2020")),("TB∕T2075—2024",("TBT2075","2024")),
    ("TB-T 3271-2011",("TBT3271","2011")),("GB T 43790-2024",("GBT43790","2024")),
    ("Q∕CR 9523-2025",("QCR9523","2025")),("QCR9523",("QCR9523","")),
    ("GB 146.2-2020",("GB146.2","2020")),("CJJ-T288-2018",("CJJT288","2018")),
])
def test_standard_variants(raw,expected): assert standards(raw)==[expected]


def test_chunk_preserves_long_text():
    text="无标点长文"*1000
    chunks=list(split_text(text,100,20))
    assert all(len(t)<=100 and t==text[s:e] for t,s,e in chunks)
    assert chunks[0][1]==0 and chunks[-1][2]==len(text)
    assert all(a[2]>=b[1] for a,b in zip(chunks,chunks[1:]))


def test_pdf_page_provenance_and_scan_report(settings,tmp_path):
    path=tmp_path/"GB 1234-2020.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page()
        page=pdf.new_page(); page.insert_text((50,50),"A useful technical passage with enough selectable text for extraction.")
        pdf.save(path)
    s=Store(settings); report=s.ingest(path)
    assert report["status"]=="imported" and len(report["warnings"])==1
    assert s.records()[0]["page"]==2 and s.records()[0]["location"]=="PDF第2页"


def test_docx_tables_positions_and_dedup(settings,tmp_path):
    path=tmp_path/"TB10009-2016.docx"; doc=Document()
    doc.add_paragraph("正文第一段。")
    table=doc.add_table(rows=1,cols=2); table.cell(0,0).text="参数"; table.cell(0,1).text="数值与单位"
    doc.add_paragraph("正文第二段。"); doc.save(path)
    s=Store(settings); a=s.ingest(path); b=s.ingest(path)
    assert a["chunks"]==3 and b["status"]=="duplicate"
    assert {r["location"] for r in s.records()}=={"段落1","段落2","表格1第1行"}
    assert all(r["page"] is None for r in s.records())


def test_corrupt_and_legacy_file_isolated(store,tmp_path):
    before=len(store.records())
    p=tmp_path/"bad.pdf"; p.write_bytes(b"broken")
    assert store.ingest(p)["status"]=="error"
    p=tmp_path/"legacy.doc"; p.write_bytes(b"legacy")
    assert store.ingest(p)["status"]=="error"
    assert len(store.records())==before


def test_exact_missing_edition_never_calls_model(store):
    client=FakeClient(); p=Pipeline(store,client)
    r=p.answer("TB/T 2073-2010 表面要求",generate=True)
    assert r["status"]=="missing_edition" and not r["sources"] and not client.calls
    assert r["available_editions"]==["2020"]


def test_multiple_editions_request_clarification(store,tmp_path):
    store.ingest(document(tmp_path,"TBT2073-2010.docx",TEXT+"历史版本。"))
    r=Pipeline(store,FakeClient()).answer("TBT2073 表面要求")
    assert r["status"]=="clarify_edition"
    r=Pipeline(store,FakeClient()).answer("TBT2073-2020 表面要求")
    assert all(s["document"]["year"]=="2020" for s in r["sources"])


def test_no_substring_standard_match(store):
    assert Pipeline(store,FakeClient()).answer("TB207 表面要求")["status"]=="missing_edition"


def test_unrelated_question_refuses_and_does_not_call(store):
    client=FakeClient(); r=Pipeline(store,client).answer("如何烹饪意大利番茄酱",generate=True)
    assert r["status"]=="insufficient_evidence" and not client.calls and not r["sources"]


def test_retrieval_only_requires_no_key(store):
    r=Pipeline(store,FakeClient()).answer("TBT2073-2020 表面裂纹")
    assert r["status"]=="retrieval_only" and r["sources"][0]["location"]=="段落1"


def test_generation_checks_quotes_and_sources(store):
    client=FakeClient(); r=Pipeline(store,client).answer("TBT2073-2020 表面裂纹",generate=True)
    assert r["status"]=="answered" and "[来源1]" in r["answer"]
    assert client.calls[0]["evidence"][0]["year"]=="2020"


@pytest.mark.parametrize("client",[FakeClient(source=99),FakeClient(quote="凭空编造的原文内容")])
def test_hallucinated_citation_blocked(store,client):
    r=Pipeline(store,client).answer("表面裂纹要求",generate=True)
    assert r["status"]=="invalid_citation" and not r["claims"]


def test_cloud_permission_blocks_remote(store):
    store.set_cloud_permission([store.documents()[0]["id"]],False)
    client=FakeClient(); r=Pipeline(store,client).answer("表面裂纹要求",generate=True)
    assert r["status"]=="cloud_permission_required" and not client.calls


def test_follow_up_resolution_and_reset():
    history=[{"question":"TB/T2073-2020 零部件表面要求","resolved_query":"TB/T2073-2020 零部件表面要求"}]
    assert "2073-2020" in resolve_query("那检验要求呢",history)[0]
    assert "TBT2073-2010" in resolve_query("那2010年版呢",history)[0]
    assert resolve_query("它的要求呢",[])[1]=="clarify_context"
    assert resolve_query("GB146.2-2020 建筑限界",history)[0]=="GB146.2-2020 建筑限界"


def test_embedding_resume_and_model_identity(store):
    client=FakeClient()
    assert store.embed_pending(client)["embedded"]==1
    assert store.embed_pending(client)["embedded"]==0
    for mode in ["dense","hybrid"]:
        assert Pipeline(store,client).answer("表面裂纹要求",mode=mode)["status"]=="retrieval_only"
    changed=Store(replace(store.settings,embedding_model="new-model"))
    r=Pipeline(changed,client).answer("表面裂纹要求",mode="hybrid")
    assert r["status"]=="api_error" and r["error_code"]=="index_not_ready"


def test_no_partial_batch_on_vector_mismatch(store):
    class Broken(FakeClient):
        def embed(self,texts): return []
    with pytest.raises(APIError): store.embed_pending(Broken())
    assert all(r["vector"] is None for r in store.records())


def test_deepseek_wire_contract(settings):
    def respond(request):
        assert str(request.url)=="https://api.deepseek.com/chat/completions"
        assert request.headers["authorization"]=="Bearer test-key"
        data=json.loads(request.content)
        assert data["model"]=="deepseek-flash" and data["response_format"]=={"type":"json_object"}
        assert data["thinking"]=={"type":"disabled"}
        return httpx.Response(200,json={"choices":[{"finish_reason":"stop","message":{"content":'{"status":"insufficient","claims":[]}'}}]})
    client=ModelClient(settings,httpx.MockTransport(respond))
    assert client.chat("JSON system","question")[0]["status"]=="insufficient"


@pytest.mark.parametrize("code",[401,429,500])
def test_api_errors_redacted(settings,code):
    client=ModelClient(settings,httpx.MockTransport(lambda r:httpx.Response(code,text="test-key secret-body")))
    with pytest.raises(APIError) as e: client.chat("s","p")
    assert e.value.code==f"http_{code}" and "test-key" not in str(e.value) and "secret-body" not in str(e.value)


def test_timeout_is_error_not_refusal(settings,store):
    def fail(r): raise httpx.ReadTimeout("secret-url")
    client=ModelClient(settings,httpx.MockTransport(fail))
    r=Pipeline(store,client).answer("表面裂纹要求",generate=True)
    assert r["status"]=="api_error" and r["error_code"]=="timeout"


@pytest.mark.parametrize("body",[{}, {"choices":[{"message":{"content":""}}]}, {"choices":[{"message":{"content":"not json"}}]}])
def test_bad_response(settings,body):
    client=ModelClient(settings,httpx.MockTransport(lambda r:httpx.Response(200,json=body)))
    with pytest.raises(APIError): client.chat("s","p")


def test_embedding_response_reorders_and_validates(settings):
    client=ModelClient(settings,httpx.MockTransport(lambda r:httpx.Response(200,json={"data":[{"index":1,"embedding":[0,1]},{"index":0,"embedding":[1,0]}]})))
    assert client.embed(["a","b"])==[[1,0],[0,1]]
    with pytest.raises(APIError): client.embed(["only one"])


def test_evaluation_matches_quote_not_just_page(store,tmp_path):
    source=store.records()[0]
    gold={"document_id":source["document_id"],"location":source["location"],"quote":TEXT[:20]}
    assert retrieval_metrics([source],[gold])["recall_at_5"]==1
    assert retrieval_metrics([source],[{**gold,"quote":"完全不同的条文内容"}])["recall_at_5"]==0
    dataset=tmp_path/"questions.jsonl"
    dataset.write_text(json.dumps({"question_id":"q1","question":"表面裂纹要求","answerable":True,"evidence_units":[gold]})+"\n")
    output=tmp_path/"results.jsonl"
    summary=run_evaluation(Pipeline(store,FakeClient()),dataset,output)
    assert summary["metrics"]["recall_at_5"]==1 and summary["questions"]==1
    assert json.loads(output.read_text())["manual_grade"]["answer_correct"] is None
    with pytest.raises(ValueError): run_evaluation(Pipeline(store,FakeClient()),dataset,output)


def test_streamlit_question_flow(store,monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("OCS_DB_PATH",str(store.settings.db_path))
    app=Path(__file__).resolve().parents[1]/"app.py"
    at=AppTest.from_file(str(app),default_timeout=15).run()
    assert not at.exception
    at.chat_input[0].set_value("TB/T 2073-2020 表面裂纹要求").run()
    assert not at.exception
    assert at.session_state["history"][-1]["status"]=="retrieval_only"
    assert any("段落1" in e.label for e in at.expander)
    at.chat_input[0].set_value("那2010年版呢").run()
    assert at.session_state["history"][-1]["status"]=="missing_edition"
    next(b for b in at.button if b.label=="清空本次对话").click().run()
    assert at.session_state["history"]==[]
