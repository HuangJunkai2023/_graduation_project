"""Fill the supplied SWJTU Attachment 9 without modifying the original form."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(__file__).resolve().parents[1]
ROOT = OUT.parent.parent
SOURCE = OUT / '开题报告内容.md'
STATS = OUT / '校验记录.json'


def target_path():
    source = SOURCE.read_text()
    student_id = re.search(r'^Student ID: (\d+)$', source, re.M)[1]
    name = re.search(r'^Student name: (.+)$', source, re.M)[1]
    return OUT / f'3.{student_id}_{name}_开题报告.docx'


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font(run, size=10.5, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')


def paragraph_format(p, size=10.5, heading=False):
    p.paragraph_format.space_before = Pt(2 if heading else 0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(size + 2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = heading
    p.paragraph_format.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # The template's document grid otherwise overrides the English line spacing.
    props = p._p.get_or_add_pPr()
    snap = props.find(qn('w:snapToGrid'))
    if snap is None:
        snap = OxmlElement('w:snapToGrid')
        props.append(snap)
    snap.set(qn('w:val'), '0')


def clear_cell(cell):
    cell._tc.clear_content()
    cell.add_paragraph()


def add_text(cell, text, size=10.5, heading=False, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph_format(p, size, heading)
    font(p.add_run(text), size, heading)
    return p


def fill_field(cell, text, size=11):
    clear_cell(cell)
    p = add_text(cell, text, size, first=True)
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def no_split(row):
    props = row._tr.get_or_add_trPr()
    if props.find(qn('w:cantSplit')) is None:
        props.append(OxmlElement('w:cantSplit'))


def add_grid(cell, lines):
    values = [[v.strip() for v in line.strip().strip('|').split('|')]
              for line in lines if not re.match(r'^\|[\s:|\-]+\|$', line)]
    table = cell.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(1.75)
    table.columns[1].width = Cm(11.2)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + name)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '777777')
        borders.append(border)
    props.append(borders)
    for i, values_row in enumerate(values):
        row = table.rows[0] if i == 0 else table.add_row()
        no_split(row)
        for j, value in enumerate(values_row):
            fill_field(row.cells[j], value, 9.5)
            row.cells[j].width = Cm(1.75 if j == 0 else 11.2)
            for p in row.cells[j].paragraphs:
                p.paragraph_format.line_spacing = Pt(11.5)
                if i == 0:
                    for run in p.runs:
                        run.bold = True
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
    # Word requires a paragraph following a nested table.
    p = cell.paragraphs[-1]
    paragraph_format(p)


def markdown(cell, text, references=False):
    clear_cell(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    lines = text.strip().splitlines()
    first = True
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            add_grid(cell, block)
            first = False
            continue
        heading = line.startswith('### ')
        text = line[4:] if heading else line
        add_text(cell, text, 9.5 if references else 10.5, heading, first)
        first = False


def page_break():
    p = OxmlElement('w:p')
    props = OxmlElement('w:pPr')
    before = OxmlElement('w:pageBreakBefore')
    props.append(before)
    keep = OxmlElement('w:keepNext')
    props.append(keep)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    props.append(snap)
    spacing = OxmlElement('w:spacing')
    for name in ['before', 'after']:
        spacing.set(qn('w:' + name), '0')
    spacing.set(qn('w:line'), '20')
    spacing.set(qn('w:lineRule'), 'exact')
    props.append(spacing)
    p.append(props)
    return p


def build():
    target = target_path()
    template = next((ROOT / '交大资料').glob('附件九*.docx'))
    before = sha256(template)
    source = SOURCE.read_text()
    fields = dict(re.findall(r'^([^\n:]+):([^\n]*)$', source.split('## ')[0], re.M))
    fields = {k.strip(): v.strip() for k, v in fields.items()}
    parts = re.split(r'^## ([^\n]+)\n', source, flags=re.M)
    sections = {parts[i]: parts[i + 1].strip() for i in range(1, len(parts), 2)}
    doc = Document(template)
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 13
    table = doc.tables[0]
    field_cells = {
        'Project title': (0, 2), 'Student name': (1, 2), 'Student ID': (1, 6),
        'Degree programme': (2, 2), 'Supervisor': (3, 3), 'Supervisor rank': (3, 7),
    }
    for name, (row, col) in field_cells.items():
        fill_field(table.cell(row, col), fields[name])
    # Preserve every option, replacing legacy Wingdings boxes with real Unicode
    # empty boxes so font substitution cannot make an option appear selected.
    assert not fields['Project type'] and not fields['Topic source'] and not fields['Guidance journal']
    choices = {
        4: [['毕业论文(理论研究)', '毕业论文(应用研究)'],
            ['毕业设计(工程设计)', '毕业设计(软件设计)']],
        5: [['教师自拟', '生产实践'], ['科研课题', '学生自拟', '其他']],
        6: [['填写', '未填写']],
    }
    for index, lines in choices.items():
        cell = table.rows[index].cells[-1]
        clear_cell(cell)
        for i, values in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            paragraph_format(p)
            if index == 6:
                font(p.add_run('毕业设计指导日志    '))
            for j, value in enumerate(values):
                if j:
                    font(p.add_run('    '))
                box = p.add_run('☐ ')
                font(box)
                box.font.name = 'DejaVu Sans'
                font(p.add_run(value))
    for row in table.rows[:7]:
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            for p in cell.paragraphs:
                paragraph_format(p)
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    if run.font.name != 'DejaVu Sans':
                        font(run)
    for index, heading in [(7, '选题意义和国内外研究现状分析'),
                           (8, '主要研究目标和研究内容'), (9, '工作计划及进度安排'),
                           (10, '参考文献')]:
        markdown(table.rows[index].cells[-1], sections[heading], references=index == 10)
        table.rows[index].height = None
        table.rows[index].height_rule = WD_ROW_HEIGHT_RULE.AUTO
    # Keep all original Chinese form items and merged-cell geometry.
    for i, row in enumerate(table.rows):
        props = row._tr.get_or_add_trPr()
        for flag in list(props.findall(qn('w:cantSplit'))):
            props.remove(flag)
        if i != 8:  # allow a long objectives section to flow rather than clip
            no_split(row)
    for i in [11, 12]:
        row = table.rows[i]
        row.height = Cm(4.8)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for p in row.cells[0].paragraphs:
            for run in p.runs:
                if run.text:
                    font(run, 10.5)
    # Extend the form by section, as explicitly allowed by the template.
    original = table._tbl
    parent = original.getparent()
    position = parent.index(original)
    groups = [(0, 8), (8, 9), (9, 10), (10, 13)]
    for i, (start, end) in enumerate(groups):
        if i:
            parent.insert(position, page_break())
            position += 1
        clone = deepcopy(original)
        for j, row in enumerate(list(clone.findall(qn('w:tr')))):
            if not start <= j < end:
                clone.remove(row)
        parent.insert(position, clone)
        position += 1
    parent.remove(original)
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph_format(p, 9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for value, field in [('第 ', None), ('', 'PAGE'), (' 页 / 共 ', None), ('', 'NUMPAGES'), (' 页', None)]:
            run = p.add_run(value)
            font(run, 9)
            if field:
                element = OxmlElement('w:fldSimple')
                element.set(qn('w:instr'), field)
                run._r.addnext(element)
    doc.core_properties.title = fields['Project title']
    doc.core_properties.author = '黄俊凯'
    doc.core_properties.subject = '西南交通大学-利兹学院本科毕业设计（论文）开题报告（附件九）'
    doc.core_properties.comments = 'Original Attachment 9 retained; opinion/signature fields intentionally blank.'
    OUT.mkdir(exist_ok=True)
    doc.save(target)
    assert sha256(template) == before
    # These blank official approval areas must not acquire generated comments.
    result = Document(target)
    assert len(result.tables) == 4
    for original_row, generated_row in [(11, 1), (12, 2)]:
        assert result.tables[-1].rows[generated_row].cells[0].text == table.rows[original_row].cells[0].text
    body = '\n'.join(sections.values())
    assert not re.search(r'\b(?:I|[Ww]e|[Mm]y|[Oo]ur)\b', body)
    assert len(re.findall(r'^\[\d+\]', sections['参考文献'], re.M)) == 7
    cited = {int(n) for n in re.findall(r'\[(\d+)\]', body.split('[1] P. Lewis')[0])}
    assert cited == set(range(1, 8)), cited
    return {
        'prepared_on': fields['Prepared on'], 'template': str(template.relative_to(ROOT)),
        'template_sha256': before, 'template_unmodified': True,
        'source_sha256': sha256(SOURCE), 'docx_sha256': sha256(target),
        'student_name': fields['Student name'], 'student_id': fields['Student ID'],
        'major': fields['Degree programme'], 'supervisor': fields['Supervisor'],
        'title': fields['Project title'], 'body_language': 'English',
        'reference_count': 7, 'research_questions': 3, 'formal_configurations': 5,
        'planned_documents': '20–30', 'planned_questions': 100,
        'development_questions': 20, 'test_questions': 80,
        'formal_executions': 400, 'planned_total_hours': 400,
        'hours_through_week24': 372, 'viva_hours': 28, 'planned_api_budget_gbp': 50,
        'blank_registration_fields': ['导师职称', '选题类型', '题目来源', '毕业设计指导日志填写情况'],
        'supervisor_opinion_blank': True, 'expert_opinion_blank': True,
        'signatures_blank': True, 'approval_dates_blank': True,
        'original_form_rows_preserved': 13, 'visual_review': False,
    }


def render(renderer, stats):
    import pymupdf
    target = target_path()
    program = Path(renderer).resolve()
    env = dict(os.environ, SAL_USE_VCLPLUGIN='svp', SAL_DISABLE_OPENCL='1', SAL_DISABLESKIA='1')
    # A locally extracted renderer is supported without a system installation.
    if program.name == 'soffice.bin' and (program.parent / 'fundamentalrc').exists():
        unpacked = program.parents[4]
        env['LD_LIBRARY_PATH'] = ':'.join([str(program.parent), str(unpacked / 'usr/lib/x86_64-linux-gnu'), env.get('LD_LIBRARY_PATH', '')])
        env['URE_BOOTSTRAP'] = 'vnd.sun.star.pathname:' + str(program.parent / 'fundamentalrc')
    with tempfile.TemporaryDirectory(prefix='ocs_swjtu_') as directory:
        command = [str(program), f'-env:UserInstallation={Path(directory).as_uri()}/profile',
                   '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(target)]
        process = subprocess.run(command, env=env, capture_output=True, text=True, timeout=60)
        if process.returncode == 81:
            process = subprocess.run(command, env=env, capture_output=True, text=True, timeout=60)
        if process.returncode:
            raise RuntimeError(process.stderr or process.stdout)
    pdf = target.with_suffix('.pdf')
    with pymupdf.open(pdf) as doc:
        text = ' '.join(' '.join(p.get_text() for p in doc).split())
        for term in [stats['student_name'], stats['student_id'], stats['supervisor'], 'Electronic and Electrical Engineering',
                     'A Study on a Large-Model-Based Intelligent Question-Answering System for the Overhead Contact System Field',
                     'RQ1', 'RQ2', 'RQ3', 'O1:', 'O5:', 'B1 BM25', 'A2 Hybrid',
                     '指导教师意见', '专家组评审意见']:
            assert term in text, term
        assert '请用英文填写' not in text
        for i, page in enumerate(doc):
            for x0, y0, x1, y1, *_ in page.get_text('blocks'):
                assert x0 >= -1 and y0 >= -1 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1, f'Page {i + 1} bounds'
        stats.update(pdf_sha256=sha256(pdf), pdf_pages=len(doc), pdf_page_bounds_checked=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--renderer', help='Path to LibreOffice soffice or soffice.bin')
    args = parser.parse_args()
    result = build()
    if args.renderer:
        render(args.renderer, result)
    STATS.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))
