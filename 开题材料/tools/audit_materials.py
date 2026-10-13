"""Local-only inventory and sample extraction audit; never calls model APIs.

Run from any directory with the project .venv Python. Original files are read-only.
Candidate status does not establish a standard's validity or permission to upload.
"""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo
import csv
import hashlib
import json
import re
import subprocess
import sys

import pymupdf as fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "开题材料/数据核查"
BASE = ROOT / "project/task1/01 data/001 规范"

# Explicit local file selection: 25 candidates, including format and version pairs.
SELECTION = [
    ("★★★★★★★★★★装备、零部件", "TB T 2073-2020 接触网零部件技术条件.pdf"),
    ("★★★★★★★★★★装备、零部件", "TB T 2074-2020 接触网零件试验方法.pdf"),
    ("★★★★★★★★★★装备、零部件", "TB∕T 2075-2024 电气化铁路接触网零部件.pdf"),
    ("★★★★★★★★★★装备、零部件", "QCR979.1～23-2023 《电气化铁路接触网零部件》.pdf"),
    ("标准/铁标", "电气化铁路接触网零部件技术条件 TBT 2073-2010.pdf"),
    ("标准/铁标", "【YY 10009】 TB 10009-2016 ★★★铁路电力牵引供电设计规范.pdf"),
    ("标准/铁标", "【YY 高规】TB 10621-2014 ★★★高速铁路设计规范.pdf"),
    ("标准/铁标", "TB T 3551-2019高速铁路牵引供电系统雷电防护技术导则.pdf"),
    ("标准/企标", "QCR 9609-2015 高速铁路电力牵引供电工程施工技术规程.pdf"),
    ("标准/企标", "QCR 9609-2015 高速铁路电力牵引供电工程施工技术规程.docx"),
    ("标准/企标", "QCR 9658-2017 客货共线铁路电力牵引用电工程施工技术规程.pdf"),
    ("标准/企标", "QCR 608-2017 ★★★电气化铁路接触网用预绞式金具.pdf"),
    ("细部设计", "QCR 9523-2018 高速铁路电力牵引供电工程细部设计和工艺质量标准.pdf"),
    ("细部设计", "Q∕CR 9523-2025 高速铁路电力牵引供电工程细部设计和工艺质量标准（有书签）.pdf"),
    ("细部设计", "Q∕CR 9523-2025 高速铁路电力牵引供电工程细部设计和工艺质量标准（报批稿）.pdf"),
    ("槽道", "TB T 3329-2013 电气化铁路接触网隧道内预埋槽道.pdf"),
    ("槽道", "TB T 3329-2013 电气化铁路接触网隧道内预埋槽道.docx"),
    ("槽道", "Q CR 1001-2023 电气化铁路接触网隧道内预埋槽道 批准20231207 实施20240229.pdf"),
    ("槽道", "Q CR 1001-2023 电气化铁路接触网隧道内预埋槽道 批准20231207 实施20240229.docx"),
    ("受电弓", "【最新 正式】GB T 43790-2024 轨道交通受流系统受电弓与接触网相互作用准则.pdf"),
    ("受电弓", "【作废】TB-T 3271-2011 轨道交通 受流系统 受电弓与接触网相互作用准则.pdf"),
    ("刚性悬挂", "TBT3252-2010电气化铁路接触网汇流排.pdf"),
    ("刚性悬挂", "城市轨道交通架空接触网技术标准-CJJ-T288-2018.pdf"),
    ("限界", "GB 146.2-2020 标准轨距铁路限界 第2部分：建筑限界.pdf"),
    ("隔离开关", "【2020-83号】工电供电函83《电气化铁路接触网隔离开关设计安装技术要求》.pdf"),
]


def standard_hint(name):
    normalized = name.upper().replace("∕", "/").replace("—", "-").replace("–", "-")
    match = re.search(r"((?:Q\s*/?\s*CR|TB(?:[ /-]*T)?|GB(?:[ /-]*T)?|CJJ[ /-]*T)\s*\d+(?:\.\d+)*(?:[～~]\d+)?)\s*-\s*(\d{4})", normalized)
    if not match:
        return "", ""
    # Canonical ID for grouping spelling variants; raw names stay in the path.
    return re.sub(r"[\s/-]+", "", match[1]), match[2]


def inspect(path):
    info = {}
    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as doc:
            # Five deterministic pages, spread through the file; not a full OCR audit.
            indices = sorted(set([0, min(1, len(doc)-1), len(doc)//4, len(doc)//2, max(0, len(doc)-2)])) if len(doc) else []
            counts = [len(re.sub(r"\s", "", doc[i].get_text())) for i in indices]
            info.update(pages=len(doc), sample_locations=";".join(str(i+1) for i in indices),
                        sample_char_counts=";".join(map(str, counts)))
            useful = sum(c >= 100 for c in counts)
            status = "text_sample_pass" if useful >= max(1, len(counts)-1) else "mixed_or_low_text" if useful else "no_useful_text_sample"
    else:
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        tables = ["\n".join("\t".join(c.text for c in r.cells) for r in t.rows) for t in doc.tables]
        info.update(pages="", sample_locations="all_paragraphs_and_tables", sample_char_counts=len("".join(paragraphs+tables)),
                    paragraph_count=len(paragraphs), table_count=len(tables))
        status = "text_sample_pass" if info["sample_char_counts"] >= 100 else "no_useful_text_sample"
    info["extraction_status"] = status
    return info


def write_csv(path, rows):
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    inventory, counts = [], {}
    for task in ["task1", "task2"]:
        files = sorted(p for p in (ROOT / "project" / task).rglob("*") if p.is_file())
        counts[task] = dict(Counter(p.suffix.lower() or "[no extension]" for p in files))
        for path in files:
            standard, year = standard_hint(path.name)
            inventory.append(dict(path=path.relative_to(ROOT).as_posix(), size_bytes=path.stat().st_size,
                                  extension=path.suffix.lower(), standard_hint=standard, year_hint=year,
                                  extraction_status="not_audited", upload_permission="unconfirmed"))
    write_csv(OUT / "全部文件清单.csv", inventory)
    candidates = []
    hashes, standards = defaultdict(list), defaultdict(list)
    for number, (directory, filename) in enumerate(SELECTION, 1):
        path = BASE / directory / filename
        row = dict(candidate_id=f"D{number:02}", path=path.relative_to(ROOT).as_posix(),
                   selection_role="comparison_only" if any(x in filename for x in ["报批稿", "作废"]) else "candidate",
                   authority_status="unverified", upload_permission="unconfirmed")
        row["standard_hint"], row["year_hint"] = standard_hint(filename)
        try:
            digest = hashlib.sha256()
            with path.open("rb") as f:
                for block in iter(lambda: f.read(1024*1024), b""): digest.update(block)
            row["sha256"] = digest.hexdigest()
            row.update(inspect(path))
            hashes[row["sha256"]].append(row["candidate_id"])
        except Exception as exc:
            row.update(extraction_status="error", error=f"{type(exc).__name__}: {exc}")
        if row["standard_hint"]:
            standards[(row["standard_hint"], row["year_hint"])].append(row["candidate_id"])
        candidates.append(row)
    for row in candidates:
        row["exact_duplicate_in_sample"] = ";".join(i for i in hashes[row.get("sha256", "")] if i != row["candidate_id"])
        row["same_standard_year_in_sample"] = ";".join(i for i in standards[(row["standard_hint"], row["year_hint"])] if i != row["candidate_id"])
    write_csv(OUT / "候选规范25份.csv", candidates)
    proc = subprocess.run([sys.executable, str(ROOT / "毕设资料/01_主要代码/selftest.py")],
                          cwd=ROOT, capture_output=True, text=True, timeout=120)
    (OUT / "参考实现离线自检.txt").write_text(proc.stdout + proc.stderr + f"\nexit_code={proc.returncode}\n", encoding="utf-8")
    summary = dict(audit_date=datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(), inventory_counts=counts,
                   candidates=len(candidates), extraction_counts=dict(Counter(r["extraction_status"] for r in candidates)),
                   baseline_selftest_exit_code=proc.returncode,
                   limitations=["PDF audit samples at most five pages, not full-document quality or OCR validation.",
                                "Filename metadata is a hint, not verification of validity or publication status.",
                                "Hash duplicate detection covers the 25 candidates only.",
                                "No embeddings, online model calls, external uploads or real QA evaluation were performed."])
    (OUT / "核查摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
