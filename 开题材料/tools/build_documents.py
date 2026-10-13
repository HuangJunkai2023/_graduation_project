"""Build editable school-template DOCX and standalone planning figures locally.

Markdown and CSV are the editable sources. Re-running regenerates derived outputs.
Does not contact providers, submit coursework or fill factual declarations.
"""
from pathlib import Path
import csv
import hashlib
import importlib.metadata
import json
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "开题材料"
SCHOOL = ROOT / "@26-27新版资料"
TITLE = "A Study on a Large-Model-Based Intelligent Question-Answering System for the Overhead Contact System Field"


def inline(p, text):
    """Preserve web links as clickable hyperlinks in the Word document."""
    for token in re.split(r"(\[[^\]]+\]\(https?://[^)]+\))", text):
        match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if match:
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), p.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            props.append(style)
            run.append(props)
            t = OxmlElement("w:t")
            t.text = match[1]
            run.append(t)
            link.append(run)
            p._p.append(link)
        else:
            p.add_run(token.replace("`", "").replace("**", ""))


def add_markdown(container, text):
    lines = text.strip().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [s.strip() for s in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", s.replace(" ", "")) for s in cells): rows.append(cells)
                i += 1
            table = container.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            for j, value in enumerate(rows[0]):
                table.rows[0].cells[j].text = value
                for run in table.rows[0].cells[j].paragraphs[0].runs: run.bold = True
            for row in rows[1:]:
                cells = table.add_row().cells
                for j, value in enumerate(row): inline(cells[j].paragraphs[0], value)
            continue
        if line.startswith("!["):
            container.add_paragraph().add_run().add_picture(str(OUT / "figures/gantt.png"), width=Inches(6.0))
        elif line.startswith("#"):
            p = container.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run(line.lstrip("# ")).bold = True
        else:
            p = container.add_paragraph()
            inline(p, line)
        i += 1


def replace_table(table, content):
    for row in list(table.rows): table._tbl.remove(row._tr)
    cell = table.add_row().cells[0]
    add_markdown(cell, content)
    return cell


def make_schedule():
    rows = list(csv.DictReader((OUT / "进度计划.csv").open(encoding="utf-8")))
    fig, ax = plt.subplots(figsize=(8.3, 6.7))
    labels = [f"{r['id']}  {r['task']}" for r in rows]
    for i, row in enumerate(rows):
        start, end = int(row["start_week"]), int(row["end_week"])
        ax.barh(i, end-start+0.8, left=start-0.4, height=0.58,
                color="#497b91" if row["id"] != "T16" else "#c5cdd0",
                edgecolor="#243c49", linewidth=0.5, hatch="//" if start >= 25 else None)
        if row["milestone"]:
            ax.plot(end+0.36, i, marker="D", markersize=4, color="#182d38")
    ax.set_yticks(range(len(rows)), labels, fontsize=7)
    ax.invert_yaxis()
    ax.set_xticks(range(5,28))
    ax.tick_params(axis="x", labelsize=7)
    ax.set_xlim(4.5,27.5)
    ax.set_xlabel("Course week (not calendar week)", fontsize=9)
    ax.set_title("ELEC3801 2026–27 | Proposed project schedule", loc="left", fontsize=12, pad=15)
    ax.grid(axis="x", alpha=0.2)
    ax.set_axisbelow(True)
    ax.axvline(24.5, color="#253744", linestyle="--", linewidth=1)
    for spine in ["top", "right", "left"]: ax.spines[spine].set_visible(False)
    fig.text(0.04, 0.03, "Diamonds = milestones. Hatched bar = post-submission viva.\nDependencies, outputs and effort are in the accompanying task table. All bars are planned work.", fontsize=8)
    fig.subplots_adjust(left=0.43, right=0.98, top=0.9, bottom=0.14)
    (OUT / "figures").mkdir(exist_ok=True)
    fig.savefig(OUT / "figures/gantt.png", dpi=220, facecolor="white")
    fig.savefig(OUT / "figures/gantt.pdf", facecolor="white")
    plt.close(fig)
    return rows


def make_outcome_map():
    handbook = Document(SCHOOL / "1.Handbook_ELEC3801_Handbook_2026-27.docx")
    table = next(t for t in handbook.tables if any("LO #" in c.text for c in t.rows[0].cells))
    proposed = [
        "检索与生成原理；报告方法部分与答辩解释", "需求清单、失败案例与研究问题", "检索实现、配置与消融实验",
        "文献阅读记录及设计选择的批判比较", "数据/成本/时间约束及范围决策记录", "离线建库与在线问答数据流及集成测试",
        "来源展示与人工复核定位、调用资源记录", "资料权限检查、真实成果与来源归属", "风险表及导师会议后的更新",
        "凭据隔离、文档指令处理与本地访问检查", "可访问界面及术语/改写问题测试", "可复现环境、解析调试和实际系统演示",
        "工具/模型/数据选择与替代方案", "甘特图、里程碑、资源与偏差记录", "报告、Review 幻灯片、视频及答辩",
        "实际周志、技术学习记录及最终个人反思",
    ]
    with (OUT / "学习成果证据规划.csv").open("w",encoding="utf-8-sig",newline="") as f:
        w = csv.writer(f)
        w.writerow(["learning_outcome", "handbook_label", "assessment_context", "planned_evidence", "actual_evidence_location", "status"])
        for i, row in enumerate(table.rows[1:]):
            values = [c.text for c in row.cells]
            w.writerow(values + [proposed[i], "", "planned_not_achieved"])
    assert len(table.rows)-1 == 16


def make_docx(tasks):
    source = (OUT / "01_Proposal_Draft_EN.md").read_text(encoding="utf-8")
    chunks = re.split(r"^## Section (\d):[^\n]*\n", source, flags=re.M)
    sections = {chunks[i]:chunks[i+1].strip() for i in range(1,len(chunks),2)}
    description, references = sections["1"].split("### References",1)
    words = re.findall(r"\S+", re.sub(r"^#+\s*", "", description, flags=re.M))
    assert len(words) <= 750, f"Section 1 exceeds 750 words: {len(words)}"
    assert len(re.findall(r"^\[\d\]",references,re.M)) == 6
    template = SCHOOL / "2.Proposal template_ELEC3801_Project_Proposal_Form.docx"
    before_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    doc = Document(template)
    assert len(doc.tables) == 10, "Unexpected template shape; inspect before editing."
    tables = list(doc.tables)
    for i in range(3): tables[0].rows[i].cells[1].text = "[TO BE COMPLETED BY STUDENT]"
    tables[0].rows[3].cells[1].text = TITLE
    doc.paragraphs[0].insert_paragraph_before("WORKING DRAFT — 13 October 2026 — Not submitted or approved").runs[0].bold = True
    replace_table(tables[1],sections["1"])
    chartcell = replace_table(tables[2],sections["2"])
    chartcell.add_paragraph("Task dependencies and effort (planned)").runs[0].bold=True
    details = chartcell.add_table(rows=1,cols=5)
    details.style="Table Grid"
    for c,v in zip(details.rows[0].cells,["ID / task","Weeks","Depends on","Hours","Output / milestone"]): c.text=v
    for r in tasks:
        values=[f"{r['id']} {r['task']}",f"{r['start_week']}–{r['end_week']}",r['depends_on'] or "—",r['planned_hours'],r['deliverable']+(" / "+r['milestone'] if r['milestone'] else "")]
        for c,v in zip(details.add_row().cells,values): c.text=v
    chartcell.add_paragraph("Planning allowance: 400 hours to final submission, plus 28 hours reserved for viva. Dependencies within the same week mean that the predecessor milestone is completed before the dependent work starts. Week 5 is an assumed start; align dates with the actual course calendar.")
    replace_table(tables[3],sections["3"])
    parts = re.split(r"^### ([^\n]+)\n",sections["4"],flags=re.M)
    professional={parts[i]:parts[i+1] for i in range(1,len(parts),2)}
    replace_table(tables[4],professional["Societal and environmental considerations"])
    replace_table(tables[5],professional["Ethics"]+"\n\n☐ No review required — student/supervisor to confirm.\n\n☐ Review required — student/supervisor to confirm.")
    replace_table(tables[6],professional["Risk Register"])
    replace_table(tables[7],professional["Security"])
    replace_table(tables[8],professional["Safety"].split("Supervisor approval,")[0])
    replace_table(tables[9],"Supervisor approval, date and means: [PENDING; TO BE RECORDED AFTER REVIEW].\n\nBefore submission, the student must obtain the supervisor's review and confirmation using the course procedure.")
    # Keep the official document's page geometry and section structure. Release
    # placeholder row heights so real content can paginate instead of clipping.
    def format_table(table):
        for row in table.rows:
            row.height=None
            row.height_rule=WD_ROW_HEIGHT_RULE.AUTO
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(5)
                    for run in p.runs:
                        run.font.name="Calibri"
                        run.font.size=Pt(10)
                for nested in cell.tables: format_table(nested)
    for t in doc.tables: format_table(t)
    doc.core_properties.title=TITLE
    doc.core_properties.subject="ELEC3801 2026–27 proposal working draft"
    doc.core_properties.author=""
    doc.core_properties.comments="Draft prepared with assistant support; personal declarations and approval remain incomplete."
    target=OUT / "Proposal_学校模板版_草稿.docx"
    doc.save(target)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before_hash
    reopened=Document(target)
    text="\n".join(node.text or "" for node in reopened.element.iter(qn("w:t")))
    assert "Remove the text below" not in text
    assert "PENDING" in text and TITLE in text
    assert len(reopened.inline_shapes)==1
    expected_urls = set(re.findall(r"\]\((https?://[^)]+)\)", references))
    actual_urls = {rel.target_ref for rel in reopened.part.rels.values() if rel.reltype == RT.HYPERLINK}
    assert len(expected_urls) == 6 and expected_urls <= actual_urls
    stats={"section_1_words_whitespace_count":len(words),"limit":750,"references":6,
           "official_template_sha256":before_hash,"template_unmodified":True,
           "docx_reopened":True,"embedded_figures":len(reopened.inline_shapes),
           "page_layout_visually_rendered":False,
           "note":"Word rendering not available in this environment. Check pagination in Word before submission; the seven-page cap applies to the final report, not this proposal."}
    (OUT / "数据核查/文档校验.json").write_text(json.dumps(stats,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(stats,ensure_ascii=False,indent=2))


def main():
    tasks=make_schedule()
    assert sum(int(r["planned_hours"]) for r in tasks if r["id"]!="T17")==400
    make_outcome_map()
    make_docx(tasks)
    packages={name:importlib.metadata.version(name) for name in ["python-docx","pymupdf","python-dotenv","numpy","matplotlib"]}
    packages["python"]=sys.version
    (OUT / "数据核查/环境版本.json").write_text(json.dumps(packages,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


if __name__=="__main__":
    main()
