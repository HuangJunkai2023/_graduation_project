"""Build the English IEEE two-column proposal requested in class.

Sources: 01_Proposal_Draft_EN.md, 进度计划.csv and 假期配置.json.
The original school forms and earlier form-based draft are retained unchanged.
"""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_documents import OUT, SCHOOL, TITLE, inline, make_schedule

TARGET = OUT / "Proposal_IEEE_English.docx"
ABSTRACT = (
    "This project will develop and evaluate a retrieval-augmented question-answering "
    "system for Chinese overhead contact system standards. It will compare lexical, "
    "dense and hybrid retrieval, test edition constraints and examine evidence-based "
    "refusal. A curated corpus and a manually checked benchmark will support reproducible "
    "evaluation of retrieval, answer correctness and source support. An initial prototype "
    "works; formal evaluation remains planned."
)


def columns(section, count):
    old=section._sectPr.find(qn("w:cols"))
    if old is not None: section._sectPr.remove(old)
    node=OxmlElement("w:cols")
    node.set(qn("w:num"),str(count))
    node.set(qn("w:space"),"720")
    section._sectPr.append(node)


def paragraph(doc, text, *, bold=False, size=10, align=None):
    p=doc.add_paragraph()
    inline(p,text)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(size * 1.1)
    p.paragraph_format.widow_control=True
    p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.name="Calibri"
        run.font.size=Pt(size)
        run.bold=bold
    # Hyperlink runs are not exposed through Paragraph.runs.
    for node in p._p.xpath('.//w:hyperlink/w:r'):
        props=node.find(qn('w:rPr'))
        fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Calibri');fonts.set(qn('w:hAnsi'),'Calibri')
        point=OxmlElement('w:sz');point.set(qn('w:val'),str(int(size*2)))
        props.extend([fonts,point])
    return p


def heading(doc,text,level=1):
    p=paragraph(doc,text,align=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before=Pt(6 if level==1 else 3)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    for run in p.runs:
        run.italic=level==2
        run.font.small_caps=level==1
    return p


def blocks(doc, text):
    """Convert source blocks into column-width prose; wide planning table is separate."""
    lines=text.strip().splitlines()
    i=0; subsection=0
    while i<len(lines):
        line=lines[i].strip()
        if not line or line.startswith('!['):i+=1;continue
        if line.startswith('### '):
            heading(doc,f"{chr(65+subsection)}. {line[4:]}",2)
            subsection+=1;i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in cells):rows.append(cells)
                i+=1
            for row in rows[1:]:
                if rows[0][0]=='ID / risk':
                    p=paragraph(doc,f"{row[0]} (likelihood / impact: {row[1]}). {row[2]}")
                else:
                    p=paragraph(doc,f"{row[0]}. {row[1]} {row[2]} {row[3]}")
                # Bold only the item identifier, not its explanation.
                text=p.text
                p.clear()
                p.add_run(row[0]+'. ').bold=True
                remainder=text[len(row[0]):].lstrip('. ')
                p.add_run(remainder)
            continue
        paragraph(doc,line)
        i+=1


def read_source():
    source=(OUT/'01_Proposal_Draft_EN.md').read_text()
    parts=re.split(r'^## Section (\d):[^\n]*\n',source,flags=re.M)
    sections={parts[i]:parts[i+1].strip() for i in range(1,len(parts),2)}
    description,references=sections['1'].split('### References',1)
    words=len(re.findall(r'\S+',re.sub(r'^#+\s*','',description,flags=re.M)))
    assert words<=750
    assert len(re.findall(r'^\[\d\]',references,re.M))==6
    assert set(re.findall(r'\[(\d)\]',description))==set('123456')
    return sections,description,references,words


def build():
    tasks=make_schedule()
    by_id={r['id']:r for r in tasks}
    for row in tasks:
        for predecessor in filter(None,row['depends_on'].split(';')):
            assert int(by_id[predecessor]['end_week'])<=int(row['start_week'])
    assert sum(int(r['planned_hours']) for r in tasks if r['id']!='T17')==400
    sections,description,references,words=read_source()
    template=SCHOOL/'3.Final report template_ELEC3801_Final_Report_Template.docx'
    template_hash=hashlib.sha256(template.read_bytes()).hexdigest()
    doc=Document(template)
    geometry=deepcopy(doc.sections[3]._sectPr)
    for name in ['headerReference','footerReference','pgNumType','type']:
        for item in list(geometry.findall(qn('w:'+name))):geometry.remove(item)
    for child in list(doc.element.body):doc.element.body.remove(child)
    doc.element.body.append(geometry)
    columns(doc.sections[0],1)
    style=doc.styles['Normal']
    style.font.name='Calibri';style.font.size=Pt(10)
    style.paragraph_format.line_spacing=Pt(11)
    style.paragraph_format.space_after=Pt(2)
    style.paragraph_format.keep_with_next=False
    style.paragraph_format.keep_together=False
    for name in ['Hyperlink','Default Paragraph Font']:
        if name in doc.styles:
            doc.styles[name].font.name='Calibri';doc.styles[name].font.size=Pt(10)

    paragraph(doc,TITLE,size=24,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'[Student name] · [SID] · [Degree programme]',size=10,align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc,'Supervisor: [Name] | ELEC3801 2026–27 | Project Proposal',size=10,align=WD_ALIGN_PARAGRAPH.CENTER)
    p=paragraph(doc,'Review version · 13 October 2026 · Personal declarations and holiday placement pending',size=9,align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after=Pt(8)
    body=doc.add_section(WD_SECTION_START.CONTINUOUS);columns(body,2)
    paragraph(doc,'Abstract — '+ABSTRACT,bold=True,size=9)
    paragraph(doc,'Index Terms — Evidence attribution, hybrid retrieval, overhead contact systems, question answering, retrieval-augmented generation.',bold=True,size=9)

    heading(doc,'I. Project Description')
    blocks(doc,description)
    heading(doc,'II. Work Plan and Milestones')
    paragraph(doc,'Fig. 1 and Table I present the weekly schedule, holiday gap, task dependencies and effort.')
    blocks(doc,sections['2'])
    heading(doc,'III. Resources and Trade-offs')
    blocks(doc,sections['3'])
    heading(doc,'IV. Professional Practice, Risk and Compliance')
    blocks(doc,sections['4'])
    heading(doc,'References')
    for line in references.strip().splitlines():
        if line.strip():
            p=paragraph(doc,line.strip(),size=9,align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent=Inches(0.18)
            p.paragraph_format.first_line_indent=Inches(-0.18)

    figure=doc.add_section(WD_SECTION_START.NEW_PAGE);columns(figure,1)
    heading(doc,'Project Schedule')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # Inline pictures need automatic line height; exact body leading clips them.
    p.paragraph_format.line_spacing=1.0
    p.add_run().add_picture(str(OUT/'figures/gantt.png'),width=Inches(7.2))
    paragraph(doc,'Fig. 1. Planned work by course week. Week 12–13 is a provisional two-week Chinese New Year break with zero planned project work; confirm its placement against the timetable. M1–M7 and MS are planned milestones.',size=8,align=WD_ALIGN_PARAGRAPH.LEFT)
    paragraph(doc,'TABLE I\nTASK DEPENDENCIES AND PLANNED EFFORT',size=8,align=WD_ALIGN_PARAGRAPH.CENTER)
    table=doc.add_table(rows=1,cols=4);table.style='Table Grid';table.autofit=False
    for cell,value in zip(table.rows[0].cells,['Task','Active course weeks','Predecessors','Hours']):cell.text=value
    config=json.loads((OUT/'假期配置.json').read_text())
    excluded={w for r in config['breaks'] for w in range(r['start_week'],r['end_week']+1)}
    activity={}
    for row in tasks:
        active=[w for w in range(int(row['start_week']),int(row['end_week'])+1) if w not in excluded]
        activity[row['id']]=active
        groups=[]
        for w in active:
            if not groups or groups[-1][-1]+1!=w:groups.append([w])
            else:groups[-1].append(w)
        label='; '.join(f'Week {g[0]}' if len(g)==1 else f'Week {g[0]}–{g[-1]}' for g in groups)
        for cell,value in zip(table.add_row().cells,[row['id'],label,row['depends_on'].replace(';',', ') or '—',row['planned_hours']]):cell.text=value
    for i,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell,width in zip(row.cells,[0.6,3.1,2.6,0.7]):
            cell.width=Inches(width)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.0
                for run in p.runs:run.font.size=Pt(8);run.bold=i==0
    # A single footer is shared across continuous and figure sections.
    footer=doc.sections[0].footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('ELEC3801 · Project Proposal · ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.name='Calibri';run.font.size=Pt(8)
    doc.core_properties.title=TITLE
    doc.core_properties.author=''
    doc.core_properties.subject='English IEEE two-column project proposal'
    doc.core_properties.comments='AI-assisted review version. Personal facts, supervisor confirmation and holiday placement remain pending.'
    doc.save(TARGET)
    assert hashlib.sha256(template.read_bytes()).hexdigest()==template_hash
    reopened=Document(TARGET)
    column_counts=[int(s._sectPr.find(qn('w:cols')).get(qn('w:num'),'1')) for s in reopened.sections]
    assert column_counts==[1,2,1]
    assert len(reopened.inline_shapes)==1
    stats={'description_words':words,'description_limit':750,'description_plus_abstract_words':words+len(ABSTRACT.split()),
           'references':6,'section_column_counts':column_counts,'layout':'Full-width title; genuine two-column English body; full-width schedule figure/table.',
           'body_font':'Calibri 10 pt, following the supplied IEEE report template',
           'template_sha256':template_hash,'template_unmodified':True,'holiday_configuration':config,
           'active_weeks':activity,'holiday_work_hours':0,'hours_through_week24':400,'viva_hours':28,
           'docx_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'visual_review':False}
    for weeks in activity.values():assert not set(weeks)&excluded
    return stats


def render(renderer,stats):
    import pymupdf
    with tempfile.TemporaryDirectory(prefix='ocs_ieee_') as tmp:
        command=[renderer,f'-env:UserInstallation={Path(tmp).as_uri()}/profile','--headless','--convert-to','pdf','--outdir',str(OUT),str(TARGET)]
        result=subprocess.run(command,capture_output=True,text=True,timeout=60)
        if result.returncode==81:result=subprocess.run(command,capture_output=True,text=True,timeout=60)
        if result.returncode:raise RuntimeError(result.stderr or result.stdout)
    pdf=TARGET.with_suffix('.pdf')
    with pymupdf.open(pdf) as doc:
        text='\n'.join(p.get_text() for p in doc)
        for expected in ['Project Description','References','TABLE I','R8','Week 12','Supervisor approval']:
            assert expected.casefold() in text.casefold(),expected
        assert not re.search('[\u4e00-\u9fff]',text),'Unexpected Chinese text in English proposal'
        figure_rects=[rect for page in doc for info in page.get_images(full=True)
                      for rect in page.get_image_rects(info[0])]
        assert any(rect.width>450 and rect.height>300 for rect in figure_rects),'Schedule figure missing or too small'
        for i,page in enumerate(doc):
            for x0,y0,x1,y1,*_ in page.get_text('blocks'):
                assert x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1, f'Content outside page {i+1}'
        stats.update(pdf_pages=len(doc),pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),pdf_text_bounds_checked=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--renderer')
    args=parser.parse_args()
    stats=build()
    if args.renderer:render(args.renderer,stats)
    (OUT/'数据核查/IEEE版校验.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in stats.items() if k not in ['active_weeks','holiday_configuration']},ensure_ascii=False,indent=2))
