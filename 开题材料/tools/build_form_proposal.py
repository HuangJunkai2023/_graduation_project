"""Build the English proposal in the supplied ELEC3801 Proposal Form/example layout.

Usage: python build_form_proposal.py [--renderer /path/to/soffice]
Keeps the official template unchanged; reads the existing Markdown and schedule.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from build_documents import OUT, SCHOOL, inline

TARGET = OUT / 'Proposal_English_Final.docx'
STATS = OUT / '数据核查/学校格式校验.json'


def format_p(p, size=10, bold=False):
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_together = True
    for r in p.runs:
        r.font.name = 'Arial'; r.font.size = Pt(size); r.bold = bold
    for r in p._p.xpath('.//w:hyperlink/w:r'):
        props=r.find(qn('w:rPr'))
        fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Arial');fonts.set(qn('w:hAnsi'),'Arial')
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));props.extend([fonts,sz])
    return p


def add_p(cell, text, size=10, bold=False):
    p=cell.add_paragraph();inline(p,text);format_p(p,size,bold)
    return p


def reset(table):
    for row in list(table.rows):table._tbl.remove(row._tr)
    cell=table.add_row().cells[0]
    format_p(cell.paragraphs[0],1)
    cell.paragraphs[0].paragraph_format.space_after=Pt(0)
    cell.paragraphs[0].paragraph_format.line_spacing=Pt(1)
    return cell


def add_grid(cell, rows, widths):
    t=cell.add_table(rows=1,cols=len(rows[0]));t.style='Table Grid';t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for i,data in enumerate(rows):
        row=t.rows[0] if i==0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for c,width,value in zip(row.cells,widths,data):
            c.width=Inches(width);c.text=value
            for p in c.paragraphs:
                format_p(p,9,i==0);p.paragraph_format.space_after=Pt(0)
    return t


def markdown(cell, text, widths=None):
    lines=text.strip().splitlines();i=0;references=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line or line.startswith('!['):continue
        if line.startswith('|'):
            table_lines=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip());i+=1
            rows=[]
            for line in table_lines:
                fields=[v.strip() for v in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',f) for f in fields):rows.append(fields)
            add_grid(cell,rows,widths or [7/len(rows[0])]*len(rows[0]));continue
        if line.startswith('### '):
            p=add_p(cell,line[4:],bold=True);p.paragraph_format.keep_with_next=True
            if line[4:]=='References':references=True
        else:add_p(cell,line,9 if references else 10)


def active_label(row, excluded):
    weeks=[w for w in range(int(row['start_week']),int(row['end_week'])+1) if w not in excluded]
    groups=[]
    for w in weeks:
        if not groups or groups[-1][-1]+1!=w:groups.append([w])
        else:groups[-1].append(w)
    return '; '.join(f'Week {g[0]}' if len(g)==1 else f'Week {g[0]}–{g[-1]}' for g in groups)


def build():
    source=(OUT/'01_Proposal_Draft_EN.md').read_text()
    parts=re.split(r'^## Section (\d):[^\n]*\n',source,flags=re.M)
    sections={parts[i]:parts[i+1].strip() for i in range(1,len(parts),2)}
    title=re.search(r'^Project title: (.+)$',source,re.M)[1]
    template=SCHOOL/'2.Proposal template_ELEC3801_Project_Proposal_Form.docx'
    example=SCHOOL/'4.Example_ELEC3801_Project_Proposal_Example.docx'
    before=hashlib.sha256(template.read_bytes()).hexdigest()
    doc=Document(template);tables=list(doc.tables);assert len(tables)==10
    # Preserve the school's cover, section order, form borders and page geometry.
    for p in doc.paragraphs:
        if 'This form must be completed' in p.text:
            for r in p.runs:r.text=r.text.replace('Week 7','Week 8 (29 October 2026)')
    cover_values=[re.search(r'^'+label+r': (.+)$',source,re.M)[1]
                  for label in ['Name and SID','Degree programme','Supervisor','Project title']]
    for i,value in enumerate(cover_values):
        tables[0].cell(i,1).text=value
        for p in tables[0].cell(i,1).paragraphs:format_p(p)
    first,second=sections['1'].split('### Reproducible method and assessment',1)
    markdown(reset(tables[1]),first)
    desc_continuation=tables[1].add_row()
    markdown(desc_continuation.cells[0],'### Reproducible method and assessment'+second)
    # Landscape chart first, as in the example; supporting task table follows.
    cell=reset(tables[2]);p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(OUT/'figures/gantt.png'),width=Inches(8.6))
    add_p(cell,'Figure 1. Planned project weeks and milestones; New Year work break: 1–14 February 2027 (Week 22–23). Table 1 gives dependencies and effort.',9)
    schedule_continuation=tables[2].add_row()
    cell=schedule_continuation.cells[0]
    markdown(cell,sections['2'])
    rows=list(csv.DictReader((OUT/'进度计划.csv').open()))
    config=json.loads((OUT/'假期配置.json').read_text())
    excluded={w for b in config['breaks'] for w in range(b['start_week'],b['end_week']+1)}
    p=add_p(cell,'Table 1. Task dependencies and planned effort',bold=True);p.paragraph_format.keep_with_next=True
    data=[['ID','Task','Active weeks','Depends on','Hours']]
    for r in rows:data.append([r['id'],r['task'],active_label(r,excluded),r['depends_on'].replace(';',', ') or '—',r['planned_hours']])
    add_grid(cell,data,[.45,4.5,2.2,2.3,.6])
    add_p(cell,'Planned effort: 372 h through Week 24 + 28 h viva = 400 h total.',9)
    markdown(reset(tables[3]),sections['3'],[1.2,2.05,1.45,2.3])
    parts=re.split(r'^### ([^\n]+)\n',sections['4'],flags=re.M)
    professional={parts[i]:parts[i+1].strip() for i in range(1,len(parts),2)}
    markdown(reset(tables[4]),professional['Societal and environmental considerations'])
    declarations=professional['Declarations for signature']
    ethics_options=declarations[declarations.index('Student ethics assessment'):declarations.index('My project work will not')].strip()
    markdown(reset(tables[5]),professional['Ethics']+'\n\n'+ethics_options)
    risk_main,risk_note=professional['Risk Register'].split('If an alternative provider',1)
    markdown(reset(tables[6]),risk_main,[1.55,.55,4.9])
    risk_continuation=tables[6].add_row()
    markdown(risk_continuation.cells[0],'If an alternative provider'+risk_note)
    markdown(reset(tables[7]),professional['Security'])
    safety=declarations[declarations.index('My project work will not'):declarations.index('Supervisor approval')].strip()
    markdown(reset(tables[8]),professional['Safety']+'\n\n'+safety,[1.4,.75,4.0,.85])
    supervisor=declarations[declarations.index('Supervisor approval'):].strip()
    markdown(reset(tables[9]),'Before submission, the supervisor should review this form, including the ethics assessment and project-specific safety assessment, and approve it below.\n\n'+supervisor)
    for table in tables:
        for row in table.rows:
            row.height=None;row.height_rule=WD_ROW_HEIGHT_RULE.AUTO
            props=row._tr.get_or_add_trPr()
            for element in list(props.findall(qn('w:cantSplit'))):props.remove(element)
    for row in [desc_continuation,schedule_continuation,risk_continuation,tables[9].rows[0]]:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    in_professional=False
    for p in doc.paragraphs:
        if p.text.startswith('Section 4:'):in_professional=True
        if not p.text.strip() and not p._p.xpath('./w:pPr/w:sectPr'):
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            if in_professional:
                p.paragraph_format.line_spacing=Pt(2)
                for run in p.runs:run.font.size=Pt(2)
        if p.text.strip() in ['Societal and environmental considerations','Ethics:','Risk Register','Security','Safety:']:
            p.paragraph_format.keep_with_next=True
    doc.core_properties.title=title;doc.core_properties.author=''
    doc.core_properties.subject='ELEC3801 2026–27 Project Proposal Form'
    doc.core_properties.comments='School Proposal Form; prepared for signature; no signature, attendance record or approval asserted.'
    doc.save(TARGET)
    assert hashlib.sha256(template.read_bytes()).hexdigest()==before
    desc=sections['1'].split('### References')[0]
    count_text='Section 1: Project Description\n'+re.sub(r'^### ','',desc,flags=re.M)
    counts={'whitespace':len(count_text.split()),'lexical':len(re.findall(r"[A-Za-z0-9]+(?:[’'-][A-Za-z0-9]+)*",count_text))}
    assert max(counts.values())<=750
    assert sum(int(r['planned_hours']) for r in rows)==400
    stats={'description_counts_including_heading':counts,'description_limit':750,'references':6,
           'hours_through_week24':372,'viva_hours':28,'total_hours':400,'template_sha256':before,
           'example_sha256':hashlib.sha256(example.read_bytes()).hexdigest(),'template_unmodified':True,
           'source_sha256':hashlib.sha256((OUT/'01_Proposal_Draft_EN.md').read_bytes()).hexdigest(),
           'docx_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'visual_review':False,
           'format':'Official Proposal Form: portrait single-column form, landscape Gantt section; Arial 10 pt body, 9 pt references and tables.',
           'proposal_deadline':'2026-10-29 (Week 8); user-confirmed classroom deadline',
           'personal_declarations':'document-only ethics option selected for review; project-specific safety assessment supplied; new assessment No selected as instructed by student; unsigned student attendance attestation and supervisor approval'}
    return stats


def render(renderer,stats):
    import pymupdf
    with tempfile.TemporaryDirectory(prefix='ocs_form_') as tmp:
        cmd=[renderer,f'-env:UserInstallation={Path(tmp).as_uri()}/profile','--headless','--convert-to','pdf','--outdir',str(OUT),str(TARGET)]
        p=subprocess.run(cmd,capture_output=True,text=True,timeout=60)
        if p.returncode==81:p=subprocess.run(cmd,capture_output=True,text=True,timeout=60)
        if p.returncode:raise RuntimeError(p.stderr or p.stdout)
    pdf=TARGET.with_suffix('.pdf')
    with pymupdf.open(pdf) as d:
        text=' '.join(' '.join(p.get_text() for p in d).split())
        supervisor_name=re.search(r'^Supervisor: (.+)$',(OUT/'01_Proposal_Draft_EN.md').read_text(),re.M)[1]
        for term in ['Section 1:', 'Section 2:', 'Section 3:', 'Section 4:', supervisor_name, '29 October 2026','R8','T17','Student signature:','Supervisor signature or confirmation']:
            assert term in text,term
        assert all(t not in text for t in ['Maya Okafor','Helen Carter','Remove the text below','IEEE two-column presentation'])
        for i,page in enumerate(d):
            for x0,y0,x1,y1,*_ in page.get_text('blocks'):
                assert x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1,f'Page {i+1} overflow'
        stats.update(pdf_pages=len(d),page_orientations=['landscape' if p.rect.width>p.rect.height else 'portrait' for p in d],pdf_bounds_checked=True,pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest())


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--renderer');args=parser.parse_args()
    stats=build()
    if args.renderer:render(args.renderer,stats)
    STATS.write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(stats,ensure_ascii=False,indent=2))
